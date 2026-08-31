#!/usr/bin/env python3
"""生成同步考勤链路说明 Excel。"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT_PATH = "/Users/jiangyongming/Project/hr/hainan/docs/钉钉接口及数据库.xlsx"


TITLE_FONT = Font(name="微软雅黑", bold=True, size=14, color="FFFFFF")
HEADER_FONT = Font(name="微软雅黑", bold=True, size=10, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
TITLE_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
ALT_FILL = PatternFill("solid", fgColor="EAF2F8")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
WRAP_ALIGN = Alignment(vertical="center", wrap_text=True)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


SUMMARY_ROWS = [
    [
        "1",
        "同步考勤组和班次",
        "topapi/attendance/getsimplegroups",
        "查询考勤组基础信息",
        "hrm_attendance_group\nhrm_attendance_group_relation_dept\nhrm_attendance_group_relation_employee",
        "写入考勤组主数据、考勤组与部门关系、考勤组与员工关系。",
        "先完整拉取远端快照，再统一替换本地快照。",
    ],
    [
        "1",
        "同步考勤组和班次",
        "topapi/attendance/shift/query",
        "查询班次详细时段",
        "hrm_attendance_shift",
        "写入班次名称、上下班时间、允许打卡窗口、休息时段等班次快照。",
        "按考勤组下挂载的每个班次逐个查询。",
    ],
    [
        "2",
        "同步考勤用户映射",
        "topapi/smartwork/hrm/employee/queryonjob",
        "查询在职员工钉钉用户 ID 列表",
        "tbattendanceuser",
        "建立钉钉用户与本地员工的映射基础数据。",
        "在职员工列表入口接口。",
    ],
    [
        "2",
        "同步考勤用户映射",
        "queryDismissionStaffIdList",
        "查询离职员工钉钉用户 ID 列表",
        "tbattendanceuser",
        "补齐离职员工的钉钉用户映射。",
        "通过阿里云 HRM SDK 调用。",
    ],
    [
        "2",
        "同步考勤用户映射",
        "topapi/smartwork/hrm/employee/v2/list",
        "按用户 ID 批量查询花名册详情",
        "tbattendanceuser",
        "补齐姓名、手机号、部门、员工 ID、考勤组等映射字段。",
        "是用户映射主数据来源。",
    ],
    [
        "2",
        "同步考勤用户映射",
        "topapi/attendance/getusergroup",
        "查询员工所属考勤组",
        "tbattendanceuser",
        "仅在本地关系缺失时，兜底补写员工当前考勤组 ID。",
        "条件触发，不是每个员工都会调用。",
    ],
    [
        "3",
        "同步考勤计划",
        "topapi/attendance/listschedule",
        "查询员工指定日期的排班计划",
        "hrm_attendance_plan",
        "写入计划 ID、员工 ID、钉钉用户 ID、考勤组 ID、班次 ID、计划打卡时间等。",
        "按日期分页拉取后覆盖目标范围内旧计划。",
    ],
    [
        "4",
        "同步打卡明细",
        "attendance/listRecord",
        "查询员工打卡记录明细",
        "hrm_attendance_clock\ntbattendancedetail",
        "一份写入业务化打卡明细，一份保留钉钉原始打卡镜像。",
        "按员工批量查询指定时间范围内的打卡记录。",
    ],
    [
        "5",
        "同步考勤报表字段",
        "topapi/attendance/getattcolumns",
        "查询考勤报表字段定义",
        "hrm_attendance_report_field",
        "写入字段 ID、字段名称、字段类型。",
        "步骤 7 回灌字段值前的字段字典。",
    ],
    [
        "6",
        "清理旧报表数据",
        "无新增钉钉接口",
        "删除目标范围内旧报表数据",
        "hrm_attendance_report_data",
        "清理当前同步员工、当前日期区间内的旧报表数据。",
        "这是本地删除步骤，不调用新的钉钉接口。",
    ],
    [
        "7a",
        "回灌日报类字段",
        "topapi/attendance/getcolumnval",
        "查询考勤日报/加班等字段值",
        "hrm_attendance_report_data",
        "写入日报类字段值，例如出勤、加班等。",
        "与步骤 7b 共用同一张报表数据表。",
    ],
    [
        "7b",
        "回灌请假类字段",
        "topapi/attendance/getleavetimebynames",
        "查询请假字段时长",
        "hrm_attendance_report_data",
        "写入请假类字段值。",
        "与步骤 7a 共用同一张报表数据表。",
    ],
]


INTERFACE_ROWS = [
    [
        "1",
        "topapi/attendance/getsimplegroups",
        "考勤组基础信息查询接口",
        "AttendanceGroupManager#fetchCurrentSnapshot",
        "同步开始时分页调用",
        "考勤组名称、考勤组类型、是否默认组、已选班次摘要、部门名称列表",
        "hrm_attendance_group\nhrm_attendance_group_relation_dept\nhrm_attendance_group_relation_employee",
        "先把远端考勤组快照完整拉回内存，再统一替换本地数据。",
    ],
    [
        "1",
        "topapi/attendance/shift/query",
        "班次详情查询接口",
        "AttendanceGroupManager#GetDetail",
        "处理每个已选班次时调用",
        "班次时段、打卡窗口、休息时段、迟到早退窗口等",
        "hrm_attendance_shift",
        "用于把班次摘要补全为可直接统计和展示的班次快照。",
    ],
    [
        "2",
        "topapi/smartwork/hrm/employee/queryonjob",
        "在职员工列表接口",
        "AttendanceUserManager#GetAndSave",
        "同步用户映射第一阶段",
        "在职员工的钉钉 userId 列表",
        "tbattendanceuser",
        "提供待同步的在职员工 userId 集合。",
    ],
    [
        "2",
        "queryDismissionStaffIdList",
        "离职员工列表接口",
        "AttendanceUserManager#GetAndSave",
        "同步用户映射第二阶段",
        "离职员工的钉钉 userId 列表",
        "tbattendanceuser",
        "补齐离职员工映射，避免历史考勤用户缺失。",
    ],
    [
        "2",
        "topapi/smartwork/hrm/employee/v2/list",
        "员工花名册详情接口",
        "AttendanceUserManager#GetUserNameByID",
        "拿到 userId 列表后批量调用",
        "姓名、手机号等花名册字段",
        "tbattendanceuser",
        "根据姓名/手机号与本地员工匹配，生成稳定映射。",
    ],
    [
        "2",
        "topapi/attendance/getusergroup",
        "员工所属考勤组查询接口",
        "AttendanceUserManager#GetGroupIDByUserID",
        "本地考勤组关系缺失时兜底调用",
        "当前员工所属考勤组 ID",
        "tbattendanceuser",
        "只做兜底补全，减少因关系快照缺失导致的 groupId 为空。",
    ],
    [
        "3",
        "topapi/attendance/listschedule",
        "排班计划查询接口",
        "AttendancePlanRecord#GetAndSaveWithoutDelete",
        "同步指定日期排班计划时调用",
        "planId、groupId、classId、classSettingId、planCheckTime、checkType、userId",
        "hrm_attendance_plan",
        "把钉钉排班计划转成本地按员工按日的计划快照。",
    ],
    [
        "4",
        "attendance/listRecord",
        "打卡记录明细接口",
        "AttendanceDetailRecord#GetAndSave",
        "同步指定时间范围打卡明细时调用",
        "打卡时间、上下班类型、打卡结果、打卡地址、设备与定位信息",
        "hrm_attendance_clock\ntbattendancedetail",
        "同一份远端记录拆成业务明细表和原始镜像表两套落库。",
    ],
    [
        "5",
        "topapi/attendance/getattcolumns",
        "考勤报表字段定义接口",
        "HrmAttendanceReportManager#UpdateReportFields",
        "同步报表字段前调用",
        "字段 ID、字段名称、字段类型",
        "hrm_attendance_report_field",
        "先刷新字段字典，再按字段查询日报和请假数据。",
    ],
    [
        "7a",
        "topapi/attendance/getcolumnval",
        "考勤日报字段值接口",
        "HrmAttendanceReportManager#UpdateAttendanceReportQuick",
        "按员工、按 15 天片段回灌日报字段时调用",
        "普通日报字段的日期和值",
        "hrm_attendance_report_data",
        "用于回写出勤、加班等日报类字段值。",
    ],
    [
        "7b",
        "topapi/attendance/getleavetimebynames",
        "请假字段时长接口",
        "HrmAttendanceReportManager#UpdateHolidayReportQuick",
        "按员工、按 15 天片段回灌请假字段时调用",
        "请假字段的日期和值",
        "hrm_attendance_report_data",
        "用于回写请假类字段值。",
    ],
]


TABLE_ROWS = [
    [
        "hrm_attendance_group",
        "考勤组主表",
        "1",
        "topapi/attendance/getsimplegroups",
        "考勤组 ID、考勤组名称、是否默认组、当前挂载班次 ID 集合",
        "保存当前钉钉考勤组快照，是班次和成员关系的主表。",
    ],
    [
        "hrm_attendance_group_relation_dept",
        "考勤组与部门关系表",
        "1",
        "topapi/attendance/getsimplegroups",
        "考勤组 ID、本地部门 ID",
        "把钉钉考勤组和本地部门关联起来。",
    ],
    [
        "hrm_attendance_group_relation_employee",
        "考勤组与员工关系表",
        "1",
        "topapi/attendance/getsimplegroups",
        "考勤组 ID、本地员工 ID",
        "记录员工当前属于哪个考勤组，用于后续用户映射和统计兜底。",
    ],
    [
        "hrm_attendance_shift",
        "考勤班次快照表",
        "1",
        "topapi/attendance/shift/query",
        "班次名称、开始结束时间、打卡窗口、休息时段、所属考勤组",
        "保存当前有效班次快照，供展示、排班和统计使用。",
    ],
    [
        "tbattendanceuser",
        "考勤用户映射表",
        "2",
        "topapi/smartwork/hrm/employee/queryonjob\nqueryDismissionStaffIdList\ntopapi/smartwork/hrm/employee/v2/list\ntopapi/attendance/getusergroup",
        "钉钉 userId、本地 empId、员工姓名、部门 ID、考勤组 ID",
        "建立钉钉考勤用户与本地员工的稳定映射关系。",
    ],
    [
        "hrm_attendance_plan",
        "考勤排班计划表",
        "3",
        "topapi/attendance/listschedule",
        "planId、groupId、classId、classSettingId、planCheckTime、checkType、workDate",
        "保存员工指定日期应执行的班次计划。",
    ],
    [
        "hrm_attendance_clock",
        "打卡业务明细表",
        "4",
        "attendance/listRecord",
        "打卡时间、上下班类型、打卡状态、打卡方式、打卡阶段、工作日期、定位信息",
        "给业务查询和统计直接使用的打卡明细表。",
    ],
    [
        "tbattendancedetail",
        "钉钉原始打卡镜像表",
        "4",
        "attendance/listRecord",
        "钉钉原始打卡记录字段，包括 groupId、checkType、timeResult、sourceType 等",
        "保留原始记录，便于问题排查和兜底取数。",
    ],
    [
        "hrm_attendance_report_field",
        "考勤报表字段表",
        "5",
        "topapi/attendance/getattcolumns",
        "字段 ID、字段名称、字段类型",
        "作为后续报表字段值回灌的字段字典。",
    ],
    [
        "hrm_attendance_report_data",
        "考勤报表数据表",
        "6、7a、7b",
        "本地删除旧数据\n topapi/attendance/getcolumnval\n topapi/attendance/getleavetimebynames",
        "员工 ID、工作日期、字段 ID、字段名称、字段值",
        "统一存放日报类字段值和请假类字段值。",
    ],
]


NON_DIRECT_ROWS = [
    [
        "hrm_attendance_date_shift",
        "按日期班次快照表",
        "不是 sync/syncAll 直接写入",
        "由考勤组配置或生效链路生成",
    ],
    [
        "hrm_emp_schedule",
        "员工排班明细表",
        "不是 sync/syncAll 直接写入",
        "由独立定时任务 AttendanceEmpScheduleTask 调 schedule/listbyusers 写入",
    ],
    [
        "hrm_attendance_history_shift",
        "历史班次表",
        "不是 sync/syncAll 直接写入",
        "来自班次维护链路，不属于同步考勤主入口",
    ],
]


CODE_ROWS = [
    [
        "入口控制器",
        "HrmAttendanceDataController#sync / #syncAll / #getSyncProgress",
        "src/main/java/com/tianye/hrsystem/controller/HrmAttendanceDataController.java",
        "同步考勤入口和进度查询入口。",
    ],
    [
        "主同步服务",
        "HrmAttendanceDataServiceImpl#SyncDataWithResume",
        "src/main/java/com/tianye/hrsystem/imple/HrmAttendanceDataServiceImpl.java",
        "按 1 到 7 步串起整条同步链路。",
    ],
    [
        "考勤组同步",
        "AttendanceGroupManager#GetAndSave",
        "src/main/java/com/tianye/hrsystem/imple/ddTalk/AttendanceGroupManager.java",
        "拉取考勤组、部门关系、员工关系和班次快照。",
    ],
    [
        "用户映射同步",
        "AttendanceUserManager#GetAndSave",
        "src/main/java/com/tianye/hrsystem/imple/ddTalk/AttendanceUserManager.java",
        "同步在职/离职员工及钉钉用户映射。",
    ],
    [
        "排班计划同步",
        "AttendancePlanRecord#GetAndSaveWithoutDelete",
        "src/main/java/com/tianye/hrsystem/imple/ddTalk/AttendancePlanRecord.java",
        "同步员工排班计划。",
    ],
    [
        "打卡明细同步",
        "AttendanceDetailRecord#GetAndSave",
        "src/main/java/com/tianye/hrsystem/imple/ddTalk/AttendanceDetailRecord.java",
        "同步打卡明细并双表落库。",
    ],
    [
        "报表字段与报表数据同步",
        "HrmAttendanceReportManager#UpdateReportFields / #UpdateAttendanceReportQuick / #UpdateHolidayReportQuick",
        "src/main/java/com/tianye/hrsystem/imple/ddTalk/HrmAttendanceReportManager.java",
        "同步报表字段字典、日报字段值和请假字段值。",
    ],
]


def apply_cell_style(cell, is_header=False, is_title=False):
    cell.border = THIN_BORDER
    if is_title:
        cell.font = TITLE_FONT
        cell.fill = TITLE_FILL
        cell.alignment = CENTER_ALIGN
    elif is_header:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
    else:
        cell.font = BODY_FONT
        cell.alignment = WRAP_ALIGN


def add_sheet(workbook, title, sheet_title, description, headers, rows, widths):
    ws = workbook.create_sheet(title)
    ws.sheet_view.showGridLines = False

    end_col = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    title_cell = ws.cell(row=1, column=1, value=sheet_title)
    apply_cell_style(title_cell, is_title=True)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    desc_cell = ws.cell(row=2, column=1, value=description)
    desc_cell.font = BODY_FONT
    desc_cell.alignment = WRAP_ALIGN
    desc_cell.fill = ALT_FILL
    desc_cell.border = THIN_BORDER

    header_row = 4
    for idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=header_row, column=idx, value=header)
        apply_cell_style(cell, is_header=True)
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row_idx, row_values in enumerate(rows, start=header_row + 1):
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            apply_cell_style(cell)
        if row_idx % 2 == 1:
            for col_idx in range(1, end_col + 1):
                ws.cell(row=row_idx, column=col_idx).fill = ALT_FILL

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{header_row + len(rows)}"

    for row_idx in range(header_row + 1, header_row + len(rows) + 1):
        ws.row_dimensions[row_idx].height = 42

    return ws


def build_workbook():
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    add_sheet(
        wb,
        "同步总览",
        "钉钉接口及数据库 - 同步总览",
        "适用范围：仅覆盖 /attendanceData/sync 与 /attendanceData/syncAll 主流程。阅读建议：先看本页总览，再看后续接口明细和落库表明细。",
        ["步骤", "同步阶段", "钉钉接口", "接口中文说明", "直接落库表", "落库说明", "备注"],
        SUMMARY_ROWS,
        [8, 18, 34, 24, 36, 32, 28],
    )
    add_sheet(
        wb,
        "接口明细",
        "钉钉接口明细",
        "按真实代码调用位置整理每个钉钉接口的用途、调用时机和直接落库去向。",
        ["步骤", "钉钉接口", "接口中文名称", "调用位置", "调用时机", "返回内容要点", "直接写入表", "说明"],
        INTERFACE_ROWS,
        [8, 34, 22, 34, 24, 34, 30, 28],
    )
    add_sheet(
        wb,
        "落库表明细",
        "落库表明细",
        "从数据库视角查看：每张表由哪个同步步骤、哪个接口写入，以及保存了什么内容。",
        ["表名", "中文名称", "来源步骤", "来源接口", "写入内容", "用途说明"],
        TABLE_ROWS,
        [34, 18, 10, 42, 40, 36],
    )
    add_sheet(
        wb,
        "非直接落库表",
        "非直接落库表",
        "这些表和考勤相关，但不属于 /attendanceData/sync 或 /attendanceData/syncAll 直接写入结果，避免排查时混淆。",
        ["表名", "中文名称", "结论", "实际来源"],
        NON_DIRECT_ROWS,
        [34, 20, 22, 44],
    )
    add_sheet(
        wb,
        "代码位置",
        "代码位置",
        "如需继续追溯实现，可从本页列出的核心类和方法入手。",
        ["模块", "类/方法", "文件路径", "说明"],
        CODE_ROWS,
        [18, 42, 68, 28],
    )
    return wb


def main():
    wb = build_workbook()
    wb.save(OUTPUT_PATH)
    print(f"已生成: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
